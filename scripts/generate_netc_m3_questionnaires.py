import os
import sys
import json
import random
import subprocess
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(SCRIPT_DIR)
DATA_DIR = os.path.join(SCRIPT_DIR, "data")
sys.path.insert(0, DATA_DIR)

from module3_data import MODULE_3_ITEMS

# Styling Palette (Professional Cisco Academic Theme)
COLOR_PRIMARY = RGBColor(0, 51, 102)     # Deep Cisco Navy
COLOR_SECONDARY = RGBColor(0, 102, 153)  # Tech Blue
COLOR_DARK = RGBColor(34, 34, 34)        # Charcoal text
COLOR_GRAY = RGBColor(100, 100, 100)     # Subtle gray
HEX_PRIMARY = "003366"
HEX_LIGHT_BG = "F4F6F9"
HEX_BORDER = "D0D7DE"

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def build_netc_doc(module_title, module_subtitle, module_code, items, output_docx_path):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Document Header
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_para.paragraph_format.space_after = Pt(2)
    run_course = header_para.add_run("NETC311: INTRODUCTION TO NETWORKS v7.0 (ITN)")
    run_course.font.name = 'Arial'
    run_course.font.size = Pt(11)
    run_course.font.bold = True
    run_course.font.color.rgb = COLOR_SECONDARY

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_after = Pt(2)
    run_title = title_para.add_run(module_title)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_para.paragraph_format.space_after = Pt(12)
    run_sub = sub_para.add_run(module_subtitle + " — Comprehensive Reviewer & Question Bank")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_GRAY

    # Student Info Box / Table
    info_table = doc.add_table(rows=2, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False

    col_widths = [Inches(4.5), Inches(2.5)]
    for row in info_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    fields = [
        ("Student Name: _________________________________", f"Date: ________________________"),
        (f"Class / Section: _______________________________", f"Score: _________ / {len(items)}")
    ]

    for row_idx, (left_txt, right_txt) in enumerate(fields):
        c_left = info_table.cell(row_idx, 0)
        c_right = info_table.cell(row_idx, 1)

        set_cell_background(c_left, HEX_LIGHT_BG)
        set_cell_background(c_right, HEX_LIGHT_BG)
        set_cell_margins(c_left, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c_right, top=80, bottom=80, left=120, right=120)

        p_left = c_left.paragraphs[0]
        p_left.paragraph_format.space_after = Pt(0)
        r = p_left.add_run(left_txt)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_DARK

        p_right = c_right.paragraphs[0]
        p_right.paragraph_format.space_after = Pt(0)
        r = p_right.add_run(right_txt)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_DARK

    set_table_borders(info_table, color="D0D7DE")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section Instructions
    inst_para = doc.add_paragraph()
    inst_para.paragraph_format.space_after = Pt(8)
    r_inst_head = inst_para.add_run("PART I: MULTIPLE CHOICE IDENTIFICATION\n")
    r_inst_head.font.name = 'Arial'
    r_inst_head.font.size = Pt(11)
    r_inst_head.font.bold = True
    r_inst_head.font.color.rgb = COLOR_PRIMARY

    r_inst_body = inst_para.add_run(
        f"General Instructions: Read each statement carefully. Select the most accurate term, acronym, or concept that "
        f"satisfies the blank. Write the letter of your choice on the blank before each number. ({len(items)} items)"
    )
    r_inst_body.font.name = 'Arial'
    r_inst_body.font.size = Pt(9.5)
    r_inst_body.font.italic = True
    r_inst_body.font.color.rgb = COLOR_DARK

    # Question Table
    q_table = doc.add_table(rows=0, cols=3)
    q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    q_table.autofit = False

    q_col_widths = [Inches(0.6), Inches(0.4), Inches(6.0)]

    processed_items = []

    for idx, item in enumerate(items, 1):
        choices = [item["a"]] + item["distractors"]
        rng = random.Random(idx * 7919)
        rng.shuffle(choices)

        letters = ["A", "B", "C", "D"]
        correct_letter = letters[choices.index(item["a"])]

        processed_items.append({
            "number": idx,
            "question": item["q"],
            "choices": list(zip(letters, choices)),
            "correct_letter": correct_letter,
            "correct_answer": item["a"],
            "topic": item["topic"],
            "explanation": item["explanation"]
        })

        row = q_table.add_row()
        for c_idx, width in enumerate(q_col_widths):
            row.cells[c_idx].width = width

        set_cell_margins(row.cells[0], top=100, bottom=100, left=80, right=80)
        set_cell_margins(row.cells[1], top=100, bottom=100, left=40, right=40)
        set_cell_margins(row.cells[2], top=100, bottom=100, left=80, right=80)

        # Col 0: Line for student's letter
        p_blank = row.cells[0].paragraphs[0]
        p_blank.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_blank.paragraph_format.space_after = Pt(0)
        r_b = p_blank.add_run("_____")
        r_b.font.name = 'Arial'
        r_b.font.size = Pt(9.5)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_DARK

        # Col 1: Number
        p_num = row.cells[1].paragraphs[0]
        p_num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_num.paragraph_format.space_after = Pt(0)
        r_n = p_num.add_run(f"{idx}.")
        r_n.font.name = 'Arial'
        r_n.font.size = Pt(9.5)
        r_n.font.bold = True
        r_n.font.color.rgb = COLOR_PRIMARY

        # Col 2: Question & Options
        p_q = row.cells[2].paragraphs[0]
        p_q.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_q.paragraph_format.space_after = Pt(2)
        r_q = p_q.add_run(item["q"])
        r_q.font.name = 'Arial'
        r_q.font.size = Pt(9.5)
        r_q.font.color.rgb = COLOR_DARK

        p_opt = row.cells[2].add_paragraph()
        p_opt.paragraph_format.space_after = Pt(4)
        p_opt.paragraph_format.left_indent = Inches(0.2)

        for l_idx, (let, ch_text) in enumerate(zip(letters, choices)):
            r_l = p_opt.add_run(f"({let}) ")
            r_l.font.name = 'Arial'
            r_l.font.size = Pt(9.0)
            r_l.font.bold = True
            r_l.font.color.rgb = COLOR_SECONDARY

            r_txt = p_opt.add_run(f"{ch_text}    ")
            r_txt.font.name = 'Arial'
            r_txt.font.size = Pt(9.0)
            r_txt.font.color.rgb = COLOR_DARK

    set_table_borders(q_table, color="E5E7EB")

    # Page Break before Answer Key
    doc.add_page_break()

    # Answer Key Section Header
    ans_header_para = doc.add_paragraph()
    ans_header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ans_header_para.paragraph_format.space_after = Pt(2)
    r_ans_h = ans_header_para.add_run("PART II: COMPLETE ANSWER KEY & TECHNICAL RATIONALES")
    r_ans_h.font.name = 'Arial'
    r_ans_h.font.size = Pt(14)
    r_ans_h.font.bold = True
    r_ans_h.font.color.rgb = COLOR_PRIMARY

    ans_sub_para = doc.add_paragraph()
    ans_sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ans_sub_para.paragraph_format.space_after = Pt(12)
    r_ans_s = ans_sub_para.add_run(f"Authoritative Cisco ITN v7.0 Reference Key — {module_title}")
    r_ans_s.font.name = 'Arial'
    r_ans_s.font.size = Pt(10)
    r_ans_s.font.italic = True
    r_ans_s.font.color.rgb = COLOR_GRAY

    # Answer Key Table
    key_table = doc.add_table(rows=1, cols=4)
    key_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    key_table.autofit = False

    key_widths = [Inches(0.6), Inches(0.8), Inches(2.2), Inches(3.4)]
    for idx, width in enumerate(key_widths):
        key_table.rows[0].cells[idx].width = width

    hdr_cells = key_table.rows[0].cells
    hdr_titles = ["Item #", "Key", "Correct Answer / Term", "Topic & Explanation"]
    for idx, (title, width) in enumerate(zip(hdr_titles, key_widths)):
        c = hdr_cells[idx]
        set_cell_background(c, HEX_PRIMARY)
        set_cell_margins(c, top=120, bottom=120, left=100, right=100)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for pitem in processed_items:
        row = key_table.add_row()
        for idx, width in enumerate(key_widths):
            row.cells[idx].width = width

        row_cells = row.cells
        set_cell_margins(row_cells[0], top=80, bottom=80, left=80, right=80)
        set_cell_margins(row_cells[1], top=80, bottom=80, left=80, right=80)
        set_cell_margins(row_cells[2], top=80, bottom=80, left=100, right=100)
        set_cell_margins(row_cells[3], top=80, bottom=80, left=100, right=100)

        if pitem["number"] % 2 == 0:
            for cell in row_cells:
                set_cell_background(cell, HEX_LIGHT_BG)

        # Col 0: Item #
        p = row_cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(pitem["number"]))
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.font.color.rgb = COLOR_DARK

        # Col 1: Correct Letter
        p = row_cells[1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"{pitem['correct_letter']}.)")
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY

        # Col 2: Term
        p = row_cells[2].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(pitem["correct_answer"])
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = COLOR_DARK

        # Col 3: Explanation & Topic
        p = row_cells[3].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)

        r_top = p.add_run(f"[{pitem['topic']}] ")
        r_top.font.name = 'Arial'
        r_top.font.size = Pt(8.5)
        r_top.font.bold = True
        r_top.font.color.rgb = COLOR_SECONDARY

        r_exp = p.add_run(pitem["explanation"])
        r_exp.font.name = 'Arial'
        r_exp.font.size = Pt(8.5)
        r_exp.font.color.rgb = COLOR_DARK

    set_table_borders(key_table, color="D0D7DE")

    doc.save(output_docx_path)
    print(f"Generated DOCX: {output_docx_path} ({len(items)} items)")

def save_json_and_md(items, base_filename, out_dir, module_title):
    json_path = os.path.join(out_dir, f"{base_filename}.json")
    json_data = {
        "title": module_title,
        "subject": "NETC311",
        "total_questions": len(items),
        "questions": []
    }
    for idx, it in enumerate(items, start=1):
        choices = [it["a"]] + it["distractors"]
        rng = random.Random(idx * 7919)
        rng.shuffle(choices)
        json_data["questions"].append({
            "number": idx,
            "question": it["q"],
            "options": choices,
            "answer": it["a"],
            "topic": it["topic"],
            "explanation": it["explanation"]
        })
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(json_data, f, indent=2)
    print(f"Saved JSON: {json_path}")

    md_path = os.path.join(out_dir, f"{base_filename}.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(f"# {module_title}\n\n")
        f.write(f"**Course:** NETC311: Introduction to Networks v7.0 (ITN)  \n")
        f.write(f"**Total Questions:** {len(items)} Items  \n\n---\n\n")
        f.write("## Part I: Identification Questionnaire\n\n")
        for idx, it in enumerate(items, start=1):
            choices = [it["a"]] + it["distractors"]
            rng = random.Random(idx * 7919)
            rng.shuffle(choices)
            f.write(f"**{idx}.** {it['q']}\n")
            for l_idx, ch in enumerate(choices):
                letter = ['A', 'B', 'C', 'D'][l_idx]
                f.write(f"- **{letter})** {ch}\n")
            f.write("\n")

        f.write("\n---\n\n## Part II: Answer Key & Detailed Rationales\n\n")
        f.write("| # | Answer | Topic | Explanation |\n")
        f.write("|---|---|---|---|\n")
        for idx, it in enumerate(items, start=1):
            f.write(f"| {idx} | **{it['a']}** | {it['topic']} | {it['explanation']} |\n")
    print(f"Saved Markdown: {md_path}")

if __name__ == "__main__":
    public_netc_dir = os.path.join(BASE_DIR, "public", "downloads", "netc311")
    docs_netc_dir = os.path.join(BASE_DIR, "docs", "downloads", "netc311")
    root_netc_dir = os.path.join(BASE_DIR, "downloads", "netc311")

    os.makedirs(public_netc_dir, exist_ok=True)
    os.makedirs(docs_netc_dir, exist_ok=True)
    os.makedirs(root_netc_dir, exist_ok=True)

    title = "Module 3: Protocols and Models"
    subtitle = "CCNA 1: Introduction to Networks v7.0"
    code = "NETC311-M3"
    doc_base = "Module 3 - Protocols and Models - Questionnaire"
    json_base = "Module_3_NotebookLM_Quiz"

    docx_path = os.path.join(public_netc_dir, f"{doc_base}.docx")
    build_netc_doc(title, subtitle, code, MODULE_3_ITEMS, docx_path)

    print(f"Converting {docx_path} to PDF via LibreOffice...")
    subprocess.run([
        'libreoffice', '--headless', '--convert-to', 'pdf',
        docx_path, '--outdir', public_netc_dir
    ], check=True)
    print("Generated PDF successfully!")

    save_json_and_md(MODULE_3_ITEMS, json_base, public_netc_dir, title)

    # Sync to docs/downloads/netc311 and downloads/netc311
    for fname in os.listdir(public_netc_dir):
        src_f = os.path.join(public_netc_dir, fname)
        shutil.copy2(src_f, os.path.join(docs_netc_dir, fname))
        shutil.copy2(src_f, os.path.join(root_netc_dir, fname))

    print("All NETC311 Module 3 questionnaire files generated and synced successfully!")
