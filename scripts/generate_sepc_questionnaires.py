import os
import sys
import json
import random
import subprocess
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(SCRIPT_DIR)
DATA_DIR = os.path.join(SCRIPT_DIR, "data")
sys.path.insert(0, DATA_DIR)

from sepc311_data import SEPC_MODULE_1_ITEMS, SEPC_MODULE_2_ITEMS

COLOR_PRIMARY = RGBColor(59, 58, 90)     # Scholarly Slate Amethyst
COLOR_SECONDARY = RGBColor(95, 93, 138)  # Ethics Muted Violet
COLOR_DARK = RGBColor(34, 34, 34)        # Charcoal text
COLOR_GRAY = RGBColor(100, 100, 100)     # Subtle gray
HEX_PRIMARY = "3B3A5A"
HEX_LIGHT_BG = "F7F6F9"
HEX_BORDER = "D1CFDC"

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D1CFDC"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def build_sepc_doc(module_title, module_subtitle, module_code, items, output_docx_path):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Document Header
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_para.paragraph_format.space_after = Pt(2)
    run_course = header_para.add_run("SEPC311: SOCIAL, ETHICAL & PROFESSIONAL ISSUES IN COMPUTING")
    run_course.font.name = 'Arial'
    run_course.font.size = Pt(11)
    run_course.font.bold = True
    run_course.font.color.rgb = COLOR_SECONDARY

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_after = Pt(2)
    run_title = title_para.add_run(module_title)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(15)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_para.paragraph_format.space_after = Pt(12)
    run_sub = sub_para.add_run(f"Comprehensive Identification & Multiple-Choice Reviewer — {len(items)} Items")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = COLOR_GRAY

    # Student Info Table
    info_table = doc.add_table(rows=1, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    info_table.columns[0].width = Inches(4.5)
    info_table.columns[1].width = Inches(2.5)

    cell_0 = info_table.cell(0, 0)
    p0 = cell_0.paragraphs[0]
    p0.paragraph_format.space_after = Pt(2)
    r_name = p0.add_run("Name: _____________________________________________")
    r_name.font.name = 'Arial'
    r_name.font.size = Pt(9.5)

    cell_1 = info_table.cell(0, 1)
    p1 = cell_1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(2)
    r_date = p1.add_run("Date: ____________ Score: _______")
    r_date.font.name = 'Arial'
    r_date.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Instructions Box
    inst_table = doc.add_table(rows=1, cols=1)
    inst_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    inst_cell = inst_table.cell(0, 0)
    set_cell_background(inst_cell, HEX_LIGHT_BG)
    set_cell_margins(inst_cell, top=100, bottom=100, left=140, right=140)
    inst_p = inst_cell.paragraphs[0]
    inst_p.paragraph_format.space_after = Pt(0)
    r_inst_lbl = inst_p.add_run("INSTRUCTIONS: ")
    r_inst_lbl.bold = True
    r_inst_lbl.font.size = Pt(9.5)
    r_inst_lbl.font.name = 'Arial'
    r_inst_lbl.font.color.rgb = COLOR_PRIMARY

    r_inst = inst_p.add_run(
        "Read each statement carefully. Identify the correct term, ethical theory, or professional tenet described in the blank (_____). "
        "Choose the letter of the correct answer from the choices provided (a, b, c, or d) and write it on the blank provided. "
        "An exhaustive Answer Key with detailed rationales is provided at the end of this document."
    )
    r_inst.font.size = Pt(9)
    r_inst.font.name = 'Arial'

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section Heading
    sec_p = doc.add_paragraph()
    sec_p.paragraph_format.space_before = Pt(10)
    sec_p.paragraph_format.space_after = Pt(6)
    r_sec = sec_p.add_run(f"PART I: IDENTIFICATION QUESTIONNAIRE ({len(items)} ITEMS)")
    r_sec.font.name = 'Arial'
    r_sec.font.size = Pt(11)
    r_sec.font.bold = True
    r_sec.font.color.rgb = COLOR_PRIMARY

    # Render items
    processed_items = []
    for idx, item in enumerate(items, start=1):
        choices = [item["a"]] + item["distractors"]
        rng = random.Random(idx * 7919)
        rng.shuffle(choices)
        correct_letter = ['a', 'b', 'c', 'd'][choices.index(item["a"])]
        processed_items.append({
            "num": idx,
            "q": item["q"],
            "choices": choices,
            "correct_letter": correct_letter,
            "correct_text": item["a"],
            "topic": item["topic"],
            "explanation": item["explanation"]
        })

        q_para = doc.add_paragraph()
        q_para.paragraph_format.space_before = Pt(4)
        q_para.paragraph_format.space_after = Pt(2)
        q_para.paragraph_format.line_spacing = 1.15

        r_num = q_para.add_run(f"_____ {idx}. ")
        r_num.bold = True
        r_num.font.name = 'Arial'
        r_num.font.size = Pt(9.5)
        r_num.font.color.rgb = COLOR_DARK

        r_text = q_para.add_run(item["q"])
        r_text.font.name = 'Arial'
        r_text.font.size = Pt(9.5)

        ch_para = doc.add_paragraph()
        ch_para.paragraph_format.left_indent = Inches(0.4)
        ch_para.paragraph_format.space_before = Pt(1)
        ch_para.paragraph_format.space_after = Pt(4)

        letters = ['a', 'b', 'c', 'd']
        for l_idx, ch in enumerate(choices):
            r_c = ch_para.add_run(f"{letters[l_idx]}.) {ch}     ")
            r_c.font.name = 'Arial'
            r_c.font.size = Pt(9)
            r_c.font.color.rgb = COLOR_DARK

    # Page Break before Answer Key
    doc.add_page_break()

    # Answer Key Heading
    ak_p = doc.add_paragraph()
    ak_p.paragraph_format.space_before = Pt(10)
    ak_p.paragraph_format.space_after = Pt(4)
    r_ak = ak_p.add_run("PART II: EXHAUSTIVE ANSWER KEY & TECHNICAL EXPLANATIONS")
    r_ak.font.name = 'Arial'
    r_ak.font.size = Pt(12)
    r_ak.font.bold = True
    r_ak.font.color.rgb = COLOR_PRIMARY

    ak_sub = doc.add_paragraph()
    ak_sub.paragraph_format.space_after = Pt(8)
    r_ak_sub = ak_sub.add_run(f"Complete answers with detailed academic explanations — {module_code}")
    r_ak_sub.font.name = 'Arial'
    r_ak_sub.font.size = Pt(9.5)
    r_ak_sub.font.color.rgb = COLOR_GRAY

    # Answer Key Table
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.1)
    table.columns[1].width = Inches(5.9)

    hdr_cells = table.rows[0].cells
    set_cell_background(hdr_cells[0], HEX_PRIMARY)
    set_cell_background(hdr_cells[1], HEX_PRIMARY)
    set_cell_margins(hdr_cells[0], top=100, bottom=100, left=100, right=100)
    set_cell_margins(hdr_cells[1], top=100, bottom=100, left=120, right=120)

    p_h0 = hdr_cells[0].paragraphs[0]
    p_h0.paragraph_format.space_after = Pt(0)
    r_h0 = p_h0.add_run("Item / Key")
    r_h0.bold = True
    r_h0.font.color.rgb = RGBColor(255, 255, 255)
    r_h0.font.size = Pt(9)
    r_h0.font.name = 'Arial'

    p_h1 = hdr_cells[1].paragraphs[0]
    p_h1.paragraph_format.space_after = Pt(0)
    r_h1 = p_h1.add_run("Topic, Correct Answer & Pedagogical Explanation")
    r_h1.bold = True
    r_h1.font.color.rgb = RGBColor(255, 255, 255)
    r_h1.font.size = Pt(9)
    r_h1.font.name = 'Arial'

    for p_item in processed_items:
        row_cells = table.add_row().cells
        set_cell_margins(row_cells[0], top=80, bottom=80, left=100, right=100)
        set_cell_margins(row_cells[1], top=80, bottom=80, left=120, right=120)

        if p_item["num"] % 2 == 0:
            set_cell_background(row_cells[0], HEX_LIGHT_BG)
            set_cell_background(row_cells[1], HEX_LIGHT_BG)

        p_c0 = row_cells[0].paragraphs[0]
        p_c0.paragraph_format.space_after = Pt(0)
        r_num_k = p_c0.add_run(f"Q{p_item['num']}: ")
        r_num_k.font.name = 'Arial'
        r_num_k.font.size = Pt(9)
        r_num_k.font.bold = True
        r_num_k.font.color.rgb = COLOR_PRIMARY

        r_let = p_c0.add_run(f"[{p_item['correct_letter'].upper()}]")
        r_let.font.name = 'Arial'
        r_let.font.size = Pt(9.5)
        r_let.font.bold = True
        r_let.font.color.rgb = COLOR_SECONDARY

        p_c1 = row_cells[1].paragraphs[0]
        p_c1.paragraph_format.space_after = Pt(1)

        r_topic = p_c1.add_run(f"Topic: {p_item['topic']}\n")
        r_topic.font.name = 'Arial'
        r_topic.font.size = Pt(8.5)
        r_topic.font.bold = True
        r_topic.font.color.rgb = COLOR_SECONDARY

        r_ans = p_c1.add_run(f"Answer: {p_item['correct_text']}\n")
        r_ans.font.name = 'Arial'
        r_ans.font.size = Pt(9)
        r_ans.font.bold = True
        r_ans.font.color.rgb = COLOR_DARK

        r_exp = p_c1.add_run(p_item["explanation"])
        r_exp.font.name = 'Arial'
        r_exp.font.size = Pt(8.5)

    set_table_borders(table)

    doc.save(output_docx_path)
    print(f"Saved DOCX: {output_docx_path}")
    return processed_items

def save_json_and_md(items, base_filename, out_dir, module_title):
    # JSON export
    json_path = os.path.join(out_dir, f"{base_filename}.json")
    json_data = {
        "title": module_title,
        "subject": "SEPC311",
        "total_questions": len(items),
        "questions": []
    }
    for idx, it in enumerate(items, start=1):
        choices = [it["a"]] + it["distractors"]
        rng = random.Random(idx * 7919)
        rng.shuffle(choices)
        json_data["questions"].append({
            "number": idx,
            "question": it["q"],
            "options": choices,
            "answer": it["a"],
            "topic": it["topic"],
            "explanation": it["explanation"]
        })
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(json_data, f, indent=2)
    print(f"Saved JSON: {json_path}")

    # Markdown export
    md_path = os.path.join(out_dir, f"{base_filename}.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(f"# {module_title}\n\n")
        f.write(f"**Course:** SEPC311: Social, Ethical, and Professional Issues in Computing  \n")
        f.write(f"**Total Questions:** {len(items)} Items  \n\n---\n\n")
        f.write("## Part I: Identification Questionnaire\n\n")
        for idx, it in enumerate(items, start=1):
            choices = [it["a"]] + it["distractors"]
            rng = random.Random(idx * 7919)
            rng.shuffle(choices)
            f.write(f"**{idx}.** {it['q']}\n")
            for l_idx, ch in enumerate(choices):
                letter = ['A', 'B', 'C', 'D'][l_idx]
                f.write(f"- **{letter})** {ch}\n")
            f.write("\n")

        f.write("\n---\n\n## Part II: Answer Key & Detailed Rationales\n\n")
        f.write("| # | Answer | Topic | Explanation |\n")
        f.write("|---|---|---|---|\n")
        for idx, it in enumerate(items, start=1):
            f.write(f"| {idx} | **{it['a']}** | {it['topic']} | {it['explanation']} |\n")
    print(f"Saved Markdown: {md_path}")

MODULE_CONFIGS = [
    {
        "title": "Module 1: Common Ethical Theories",
        "subtitle": "SEPC311 — Social, Ethical & Professional Issues in Computing",
        "code": "SEPC311-M1",
        "filename_base": "Module 1 - Common Ethical Theories - Questionnaire",
        "json_md_base": "SEPC311_Module_1_Quiz",
        "items": SEPC_MODULE_1_ITEMS
    },
    {
        "title": "Module 2: Computer Ethics and Professional Codes",
        "subtitle": "SEPC311 — Social, Ethical & Professional Issues in Computing",
        "code": "SEPC311-M2",
        "filename_base": "Module 2 - Computer Ethics and Professional Codes - Questionnaire",
        "json_md_base": "SEPC311_Module_2_Quiz",
        "items": SEPC_MODULE_2_ITEMS
    }
]

# Ensure download directories exist
PUBLIC_DL_DIR = os.path.join(BASE_DIR, "public", "downloads", "sepc311")
DOCS_DL_DIR = os.path.join(BASE_DIR, "docs", "downloads", "sepc311")
os.makedirs(PUBLIC_DL_DIR, exist_ok=True)
os.makedirs(DOCS_DL_DIR, exist_ok=True)

for config in MODULE_CONFIGS:
    docx_path = os.path.join(PUBLIC_DL_DIR, f"{config['filename_base']}.docx")
    build_sepc_doc(
        config["title"],
        config["subtitle"],
        config["code"],
        config["items"],
        docx_path
    )
    print(f"Converting {docx_path} to PDF via LibreOffice...")
    subprocess.run([
        'libreoffice', '--headless', '--convert-to', 'pdf',
        docx_path, '--outdir', PUBLIC_DL_DIR
    ], check=True)
    print(f"Generated PDF for {config['code']} successfully!")

    save_json_and_md(config["items"], config["json_md_base"], PUBLIC_DL_DIR, config["title"])

# Sync public/downloads/sepc311 to docs/downloads/sepc311
for f in os.listdir(PUBLIC_DL_DIR):
    shutil.copy2(os.path.join(PUBLIC_DL_DIR, f), os.path.join(DOCS_DL_DIR, f))

print(f"All SEPC311 documents successfully generated and synced to docs/downloads/sepc311/!")
